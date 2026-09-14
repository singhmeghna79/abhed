#!/usr/bin/env python
"""
Build the Abhed (Zybuu) investor/customer deck.

Source of truth for every claim in this deck:
  - /Users/yuvrajsingh/titan/docs/vision.md
  - /Users/yuvrajsingh/titan/web/zybuu/index.html
  - /Users/yuvrajsingh/titan/web/zybuu/abhed/index.html
  - /Users/yuvrajsingh/titan/docs/guide/README.md

No invented numbers, customers, revenue or certifications. Zybuu holds no
SOC 2 / ISO 27001 / HIPAA certification (stated plainly on slide 9). Pricing
is a proposal only and is labelled "proposed" everywhere it appears.
"""

from pptx import os
import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
import copy

# ---------------------------------------------------------------- palette --
BG        = RGBColor(0x06, 0x09, 0x0F)
SURFACE   = RGBColor(0x0D, 0x13, 0x1C)
BORDER    = RGBColor(0x18, 0x22, 0x31)
TEXT      = RGBColor(0xE8, 0xEE, 0xF7)
SECONDARY = RGBColor(0xB0, 0xBF, 0xD2)
MUTED     = RGBColor(0x7A, 0x8A, 0xA0)
ACCENT    = RGBColor(0x3B, 0xA9, 0xFF)   # neon blue
ACCENT2   = RGBColor(0x8B, 0x6C, 0xFF)   # violet
GREEN     = RGBColor(0x3D, 0xD6, 0x8C)   # "live"
AMBER     = RGBColor(0xE0, 0x8A, 0x4C)   # warn / honest slide
RED_SOFT  = RGBColor(0xFF, 0x5F, 0x57)

HEAD_FONT = "Helvetica Neue"
MONO_FONT = "Menlo"

# ---------------------------------------------------------------- geometry --
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN  = Inches(0.6)
CONTENT_W = SLIDE_W - 2 * MARGIN

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]

SLIDE_NO = {"n": 0}


# ------------------------------------------------------------------ helpers
def add_slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    bg.shadow.inherit = False
    # send to back
    sp = bg._element
    sp.getparent().remove(sp)
    s.shapes._spTree.insert(2, sp)
    SLIDE_NO["n"] += 1
    return s


def set_font(run, size=18, color=TEXT, bold=False, font=HEAD_FONT, italic=False):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font


def no_autosize(tf):
    from pptx.enum.text import MSO_AUTO_SIZE
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE


def add_text(slide, left, top, width, height, text, size=18, color=TEXT,
             bold=False, font=HEAD_FONT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             italic=False, line_spacing=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    no_autosize(tf)
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        r = p.add_run()
        r.text = line
        set_font(r, size, color, bold, font, italic)
    return box


def add_multirun_para(tf, runs, align=PP_ALIGN.LEFT, new_para=True, line_spacing=None, space_after=None):
    p = tf.paragraphs[0] if (not new_para and len(tf.paragraphs) == 1 and not tf.paragraphs[0].runs) else tf.add_paragraph()
    p.alignment = align
    if line_spacing:
        p.line_spacing = line_spacing
    if space_after is not None:
        p.space_after = Pt(space_after)
    for (text, size, color, bold, font, italic) in runs:
        r = p.add_run()
        r.text = text
        set_font(r, size, color, bold, font, italic)
    return p


def rounded_rect(slide, left, top, width, height, fill=SURFACE, line_color=BORDER,
                  line_w=Pt(1), radius=0.06, shadow=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    try:
        shp.adjustments[0] = radius
    except Exception:
        pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line_color is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line_color
        shp.line.width = line_w
    shp.shadow.inherit = False
    return shp


def rect(slide, left, top, width, height, fill=SURFACE, line_color=None, line_w=Pt(1)):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line_color is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line_color
        shp.line.width = line_w
    shp.shadow.inherit = False
    return shp


def thin_rule(slide, left, top, width, color=ACCENT, height=Pt(2.2)):
    return rect(slide, left, top, width, height, fill=color, line_color=None)


def add_footer(slide, page_num):
    y = SLIDE_H - Inches(0.42)
    add_text(slide, MARGIN, y, Inches(4), Inches(0.3),
              "Zybuu · Abhed", size=10.5, color=MUTED, font=MONO_FONT)
    add_text(slide, SLIDE_W - MARGIN - Inches(1.5), y, Inches(1.5), Inches(0.3),
              f"{page_num:02d} / 12", size=10.5, color=MUTED, font=MONO_FONT, align=PP_ALIGN.RIGHT)


def kicker(slide, text, top=MARGIN):
    # small accent rule + uppercase mono label, matches site's .kicker
    rule = thin_rule(slide, MARGIN, top + Pt(6), Inches(0.28), color=ACCENT, height=Pt(2.4))
    add_text(slide, MARGIN + Inches(0.36), top, Inches(9), Inches(0.32),
              text.upper(), size=12.5, color=ACCENT, bold=True, font=MONO_FONT)


def headline(slide, text, top, size=40, color=TEXT, width=None, align=PP_ALIGN.LEFT):
    w = width or (CONTENT_W)
    box = add_text(slide, MARGIN, top, w, Inches(2.0), text, size=size, color=color,
                    bold=True, font=HEAD_FONT, align=align, line_spacing=1.02)
    return box


def pill(slide, left, top, width, height, text, fill, text_color, size=11):
    shp = rounded_rect(slide, left, top, width, height, fill=fill, line_color=None, radius=0.5)
    tf = shp.text_frame
    no_autosize(tf)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Pt(4)
    tf.margin_right = Pt(4)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    set_font(r, size, text_color, True, MONO_FONT)
    return shp


def card(slide, left, top, width, height, title, body, title_size=15, body_size=12.5,
          title_color=TEXT, body_color=SECONDARY, accent_bar=None, pad=Inches(0.2)):
    c = rounded_rect(slide, left, top, width, height, fill=SURFACE, line_color=BORDER, radius=0.07)
    if accent_bar:
        thin_rule(slide, left, top, width, color=accent_bar, height=Pt(2.6))
    tx_top = top + pad
    box = slide.shapes.add_textbox(left + pad, tx_top, width - 2 * pad, height - 2 * pad)
    tf = box.text_frame
    no_autosize(tf)
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    set_font(r, title_size, title_color, True, HEAD_FONT)
    p.space_after = Pt(6)
    if body:
        p2 = tf.add_paragraph()
        p2.line_spacing = 1.18
        r2 = p2.add_run()
        r2.text = body
        set_font(r2, body_size, body_color, False, HEAD_FONT)
    return c


def term_window(slide, left, top, width, height, title="~/demo — abhed"):
    win = rounded_rect(slide, left, top, width, height, fill=RGBColor(0x0B, 0x0E, 0x14),
                        line_color=RGBColor(0x23, 0x2B, 0x38), radius=0.045)
    bar_h = Inches(0.34)
    bar = rect(slide, left, top, width, bar_h, fill=RGBColor(0x0F, 0x14, 0x1C), line_color=None)
    # traffic dots
    dot_r = Inches(0.075)
    dot_y = top + bar_h / 2 - dot_r / 2
    colors = [RGBColor(0xFF, 0x5F, 0x57), RGBColor(0xFE, 0xBC, 0x2E), RGBColor(0x28, 0xC8, 0x40)]
    for i, col in enumerate(colors):
        dx = left + Inches(0.14) + i * Inches(0.19)
        d = slide.shapes.add_shape(MSO_SHAPE.OVAL, dx, dot_y, dot_r, dot_r)
        d.fill.solid(); d.fill.fore_color.rgb = col
        d.line.fill.background()
        d.shadow.inherit = False
    add_text(slide, left + Inches(0.85), top, width - Inches(2.6), bar_h, title,
              size=10, color=RGBColor(0x8A, 0x96, 0xA8), font=MONO_FONT, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, left + width - Inches(1.85), top, Inches(1.65), bar_h, "recorded session",
              size=8.5, color=RGBColor(0x8A, 0x96, 0xA8), font=MONO_FONT, anchor=MSO_ANCHOR.MIDDLE,
              align=PP_ALIGN.RIGHT)
    body_top = top + bar_h
    body_h = height - bar_h
    return body_top, body_h


TERM_CMD  = RGBColor(0xEE, 0xF3, 0xFA)
TERM_TOOL = RGBColor(0x6F, 0xB3, 0xEE)
TERM_RES  = RGBColor(0x8A, 0x96, 0xA8)
TERM_OK   = RGBColor(0x3F, 0xD6, 0x8C)
TERM_WARN = RGBColor(0xE0, 0xB0, 0x4C)
TERM_DIM  = RGBColor(0x77, 0x83, 0x9A)


def term_lines(slide, left, top, width, height, lines, size=12.5):
    box = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.12), width - Inches(0.4), height - Inches(0.2))
    tf = box.text_frame
    no_autosize(tf)
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    first = True
    for kind, text in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.line_spacing = 1.32
        r = p.add_run()
        r.text = text
        color = {"cmd": TERM_CMD, "tool": TERM_TOOL, "res": TERM_RES,
                 "ok": TERM_OK, "warn": TERM_WARN, "dim": TERM_DIM}.get(kind, TERM_RES)
        bold = kind == "cmd"
        set_font(r, size, color, bold, MONO_FONT)
    return box


def arrow(slide, x1, y1, x2, y2, color=ACCENT, width=Pt(2)):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    conn.line.color.rgb = color
    conn.line.width = width
    line = conn.line._get_or_add_ln()
    tail = line.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'})
    line.append(tail)
    conn.shadow.inherit = False
    return conn


# =============================================================================
# SLIDE 1 — Title
# =============================================================================
s = add_slide()

# subtle accent glow bar top
thin_rule(s, MARGIN, Inches(1.55), Inches(1.0), color=ACCENT, height=Pt(3))

add_text(s, MARGIN, Inches(0.9), Inches(6), Inches(0.4), "ZYBUU", size=16,
          color=ACCENT, bold=True, font=MONO_FONT)

headline(s, "Abhed", Inches(1.9), size=64, color=TEXT)
add_text(s, MARGIN, Inches(3.0), Inches(11.5), Inches(1.6),
          "The agent harness for work that\ncannot leave the building.",
          size=32, color=SECONDARY, bold=True, font=HEAD_FONT, line_spacing=1.08)

add_text(s, MARGIN, Inches(5.35), Inches(8), Inches(0.4),
          "Zybuu  ·  September 2026", size=15, color=SECONDARY, font=HEAD_FONT)
add_text(s, MARGIN, Inches(5.78), Inches(8), Inches(0.4),
          "Confidential — shared on request.", size=12.5, color=MUTED, font=MONO_FONT, italic=True)

add_footer(s, 1)

# =============================================================================
# SLIDE 2 — The thesis
# =============================================================================
s = add_slide()
kicker(s, "The thesis")
headline(s, "Every enterprise will run agents.", Inches(1.15), size=38)
headline(s, "Most cannot run them the way the market sells them.", Inches(1.95), size=38, color=ACCENT)

add_text(s, MARGIN, Inches(3.15), Inches(11.6), Inches(1.3),
          "The market is selling agents one way: a vendor's cloud, the vendor's model, "
          "on the vendor's terms — with the organisation's data leaving the building on every turn.",
          size=17, color=SECONDARY, line_spacing=1.3)

add_text(s, MARGIN, Inches(4.15), Inches(11.6), Inches(1.5),
          "Banks and insurers, hospitals, defence and government, utilities, and every company with a "
          "data-residency clause are offered a choice: wait, accept a weaker product, or build the whole "
          "thing themselves. Most are building. Most of what they build is the same thing, badly.",
          size=15, color=MUTED, line_spacing=1.35)

card(s, MARGIN, Inches(5.85), CONTENT_W, Inches(0.95),
     "Zybuu builds the infrastructure that lets an organisation run AI workloads under its own control",
     "SaaS, privately hosted, or fully air-gapped — with the same guarantees at every tier.",
     title_size=15.5, body_size=13, title_color=TEXT, accent_bar=ACCENT)

add_footer(s, 2)

# =============================================================================
# SLIDE 3 — The problem in one picture
# =============================================================================
s = add_slide()
kicker(s, "The problem, in one picture")
headline(s, "Data leaves. Or the model comes to the data.", Inches(1.1), size=32)

diag_top = Inches(2.3)
diag_h = Inches(4.35)
col_w = Inches(5.6)
gap = Inches(0.5)
left_x = MARGIN
right_x = MARGIN + col_w + gap

# ---- Left: cloud-agent flow ----
rounded_rect(s, left_x, diag_top, col_w, diag_h, fill=SURFACE, line_color=BORDER, radius=0.04)
add_text(s, left_x + Inches(0.25), diag_top + Inches(0.2), col_w - Inches(0.5), Inches(0.4),
          "CLOUD-AGENT FLOW", size=13, color=AMBER, bold=True, font=MONO_FONT)

def box_label(slide, x, y, w, h, text, fill=SURFACE, border=BORDER, tcolor=TEXT, size=12.5, bold=False):
    b = rounded_rect(slide, x, y, w, h, fill=fill, line_color=border, radius=0.12)
    tf = b.text_frame
    no_autosize(tf); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Pt(6); tf.margin_right = Pt(6)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    set_font(r, size, tcolor, bold, HEAD_FONT)
    return b

bw, bh = Inches(2.3), Inches(0.62)
bx = left_x + (col_w - bw) / 2
y0 = diag_top + Inches(0.85)
box_label(s, bx, y0, bw, bh, "Your data", fill=RGBColor(0x14,0x1B,0x27), border=BORDER)
y1 = y0 + bh + Inches(0.55)
box_label(s, bx, y1, bw, bh, "Vendor's cloud", fill=RGBColor(0x2A, 0x1A, 0x10), border=AMBER, tcolor=AMBER)
y2 = y1 + bh + Inches(0.55)
box_label(s, bx, y2, bw, bh, "Vendor's model", fill=RGBColor(0x14,0x1B,0x27), border=BORDER)
arrow(s, bx + bw/2, y0 + bh, bx + bw/2, y1, color=AMBER)
arrow(s, bx + bw/2, y1 + bh, bx + bw/2, y2, color=AMBER)
add_text(s, left_x + Inches(0.25), diag_top + diag_h - Inches(0.62), col_w - Inches(0.5), Inches(0.45),
          "Data leaves the building on every turn.", size=12, color=SECONDARY, italic=True, align=PP_ALIGN.CENTER)

# ---- Right: Abhed flow ----
rounded_rect(s, right_x, diag_top, col_w, diag_h, fill=SURFACE, line_color=ACCENT, radius=0.04, line_w=Pt(1.4))
add_text(s, right_x + Inches(0.25), diag_top + Inches(0.2), col_w - Inches(0.5), Inches(0.4),
          "ABHED", size=13, color=ACCENT, bold=True, font=MONO_FONT)

bx2 = right_x + (col_w - bw) / 2
box_label(s, bx2, y0, bw, bh, "Your data", fill=RGBColor(0x14,0x1B,0x27), border=BORDER)
box_label(s, bx2, y1, bw, bh, "Abhed (your boundary)", fill=RGBColor(0x0B,0x25,0x40), border=ACCENT, tcolor=ACCENT)
box_label(s, bx2, y2, bw, bh, "Any model — incl. yours", fill=RGBColor(0x14,0x1B,0x27), border=BORDER)
arrow(s, bx2 + bw/2, y1, bx2 + bw/2, y0 + bh, color=ACCENT)   # model comes to data (down arrow reversed feel)
arrow(s, bx2 + bw/2, y2, bx2 + bw/2, y1 + bh, color=ACCENT)
add_text(s, right_x + Inches(0.25), diag_top + diag_h - Inches(0.62), col_w - Inches(0.5), Inches(0.45),
          "The model comes to the data. Nothing leaves by default.", size=12, color=SECONDARY, italic=True, align=PP_ALIGN.CENTER)

add_footer(s, 3)

# =============================================================================
# SLIDE 4 — What Zybuu builds
# =============================================================================
s = add_slide()
kicker(s, "What Zybuu builds")
headline(s, "The layer underneath the agent.", Inches(1.1), size=36)
add_text(s, MARGIN, Inches(1.95), Inches(11.6), Inches(0.9),
          "Not a better chatbot: the part that decides what a model may do, does it safely, "
          "records it, and can prove afterwards what happened.",
          size=15.5, color=SECONDARY, line_spacing=1.3)

tiers = [
    ("SaaS", "Hosted by Zybuu. The fastest path to running an agent under policy."),
    ("Private", "Deployed in the customer's own cloud or data centre."),
    ("Air-gapped", "A signed offline bundle, for the rack with no route out."),
]
tw = Inches(3.65)
gap2 = Inches(0.32)
tx0 = MARGIN
ty = Inches(3.15)
th = Inches(2.0)
for i, (title, body) in enumerate(tiers):
    x = tx0 + i * (tw + gap2)
    card(s, x, ty, tw, th, title, body, title_size=19, body_size=13.5, accent_bar=ACCENT)

add_text(s, MARGIN, ty + th + Inches(0.35), CONTENT_W, Inches(0.6),
          "Same guarantees at every tier.", size=20, color=ACCENT, bold=True, font=HEAD_FONT)
add_text(s, MARGIN, ty + th + Inches(0.95), CONTENT_W, Inches(0.6),
          "Abhed is the first product on this layer.", size=13.5, color=MUTED, font=HEAD_FONT, italic=True)

add_footer(s, 4)

# =============================================================================
# SLIDE 5 — Under your control
# =============================================================================
s = add_slide()
kicker(s, "Abhed")
headline(s, "Five things “under your control” means.", Inches(1.1), size=32)
add_text(s, MARGIN, Inches(1.85), Inches(11.6), Inches(0.5),
          "Each one is a property this repository tests, not a line of copy.",
          size=13.5, color=MUTED, italic=True)

items5 = [
    ("Runs where the data is", "One static binary, from a laptop to an air-gapped rack, with a signed offline bundle for the rack."),
    ("Runs any model", "Twenty providers over three wire formats, including the one on your own GPUs. Changing vendors is a line of config."),
    ("Sandboxed by default", "Inside a boundary the operator sets and the agent cannot lift, with an approval policy the operator owns."),
    ("Every action on the record", "An append-only event log, replayable step by step, with every approval and refusal and who made it."),
    ("Embeds without weakening", "The SDK gives a program the same loop with the same guarantees. The console is a reference app, not the product."),
]
cw = Inches(2.28)
cgap = Inches(0.14)
cy = Inches(2.55)
chh = Inches(3.9)
cx0 = MARGIN
for i, (title, body) in enumerate(items5):
    x = cx0 + i * (cw + cgap)
    card(s, x, cy, cw, chh, title, body, title_size=13.5, body_size=11, accent_bar=(ACCENT if i % 2 == 0 else ACCENT2))

add_footer(s, 5)

# =============================================================================
# SLIDE 6 — Live demo
# =============================================================================
s = add_slide()
kicker(s, "Live demo")
headline(s, "A real run, captured.", Inches(1.1), size=34)

term_left = MARGIN
term_top = Inches(1.95)
term_w = CONTENT_W
term_h = Inches(4.55)
body_top, body_h = term_window(s, term_left, term_top, term_w, term_h)

lines = [
    ("cmd",  "$ abhed doctor"),
    ("ok",   "checking endpoint... ok"),
    ("ok",   "checking tool calling... ok"),
    ("ok",   "Ready."),
    ("dim",  ""),
    ("cmd",  '$ abhed -p "Add a FToC function ... and a table-driven test ... Then run go test."'),
    ("tool", "● glob tempconv*"),
    ("tool", "● read tempconv.go"),
    ("warn", "│ path must be absolute. Did you mean ~/demo/tempconv.go?"),
    ("tool", "● edit ~/demo/tempconv.go"),
    ("tool", "● write ~/demo/tempconv_test.go"),
    ("tool", "● bash Run tests"),
    ("ok",   "└ exit 0"),
    ("dim",  ""),
    ("dim",  "9 turns · 30,193 in / 1,825 out tokens"),
]
term_lines(s, term_left, body_top, term_w, body_h, lines, size=13)

add_text(s, MARGIN, Inches(6.65), CONTENT_W, Inches(0.5),
          "26B model, local, nothing left the machine.", size=15, color=SECONDARY, italic=True, font=HEAD_FONT)

add_footer(s, 6)

# =============================================================================
# SLIDE 7 — How it competes
# =============================================================================
s = add_slide()
kicker(s, "How it competes")
headline(s, "We don't compete on the model.", Inches(1.05), size=30)
add_text(s, MARGIN, Inches(1.72), Inches(11.6), Inches(0.45),
          "We run theirs, or yours. We compete on where the agent runs, what it can prove, and what it costs to operate.",
          size=13, color=MUTED, italic=True)

cols7 = [
    ("Cloud coding agents", "Claude Code, OpenAI Codex CLI",
     [("Built for", "One developer, one vendor's models"),
      ("Where model runs", "The vendor's cloud"),
      ("Abhed's difference", "Runs where the data is, on any model, with policy and a record the organisation owns.")],
     False),
    ("Terminal harnesses", "pi.dev",
     [("Built for", "One developer, many models, in a terminal — better than Abhed for that developer"),
      ("Where model runs", "Wherever you point it"),
      ("Abhed's difference", "Built for a platform team serving an organisation: users, tenants, policy, audit, a server, an SDK.")],
     False),
    ("Agent frameworks", "CrewAI and frameworks like it",
     [("Built for", "A Python team assembling agents"),
      ("Left to you", "The sandbox, the approvals, the audit log, the deployment"),
      ("Abhed's difference", "The assembled, hardened thing — with extensions, skills and MCP underneath for the parts that should be yours.")],
     False),
    ("Abhed", None,
     [("Built for", "Organisations whose data cannot leave"),
      ("Where it runs", "Laptop, rack, or air-gapped enclave — one binary"),
      ("Model", "Twenty providers, including the one on your GPUs"),
      ("What you own", "Every action, every decision, every byte — recorded and replayable.")],
     True),
]
colw = Inches(3.03)
colgap = Inches(0.14)
coly = Inches(2.35)
colh = Inches(4.35)
colx0 = MARGIN
for i, (title, sub, rows, is_us) in enumerate(cols7):
    x = colx0 + i * (colw + colgap)
    fill = RGBColor(0x0B, 0x25, 0x40) if is_us else SURFACE
    border = ACCENT if is_us else BORDER
    c = rounded_rect(s, x, coly, colw, colh, fill=fill, line_color=border,
                      line_w=Pt(1.6 if is_us else 1), radius=0.05)
    box = s.shapes.add_textbox(x + Inches(0.15), coly + Inches(0.15), colw - Inches(0.3), colh - Inches(0.3))
    tf = box.text_frame
    no_autosize(tf); tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = title
    set_font(r, 15, ACCENT if is_us else TEXT, True, HEAD_FONT)
    p.space_after = Pt(2)
    if sub:
        p2 = tf.add_paragraph()
        r2 = p2.add_run(); r2.text = sub
        set_font(r2, 9.5, MUTED, False, MONO_FONT, italic=True)
        p2.space_after = Pt(8)
    else:
        p.space_after = Pt(10)
    for label, val in rows:
        pl = tf.add_paragraph()
        pl.space_before = Pt(7)
        rl = pl.add_run(); rl.text = label.upper()
        set_font(rl, 8.5, ACCENT if is_us else MUTED, True, MONO_FONT)
        pv = tf.add_paragraph()
        pv.line_spacing = 1.15
        rv = pv.add_run(); rv.text = val
        set_font(rv, 10.5, TEXT if is_us else SECONDARY, False, HEAD_FONT)

add_text(s, MARGIN, Inches(6.85), CONTENT_W, Inches(0.35),
          "Descriptions of other products are from their public documentation and meant fairly.",
          size=10, color=MUTED, italic=True)

add_footer(s, 7)

# =============================================================================
# SLIDE 8 — Evidence
# =============================================================================
s = add_slide()
kicker(s, "Evidence")
headline(s, "Numbers that survive a technical diligence call.", Inches(1.1), size=32)

figs = [
    ("24 / 24", "adversarial attacks blocked\nin the red-team suite"),
    ("584", "test functions\nacross the engine"),
    ("20", "model providers\nbehind one abstraction"),
    ("0", "bytes leaving the network\nby default"),
]
fw = Inches(2.75)
fgap = Inches(0.22)
fy = Inches(2.6)
fh = Inches(2.55)
fx0 = MARGIN + (CONTENT_W - (4 * fw + 3 * fgap)) / 2
for i, (num, cap) in enumerate(figs):
    x = fx0 + i * (fw + fgap)
    rounded_rect(s, x, fy, fw, fh, fill=SURFACE, line_color=BORDER, radius=0.07)
    add_text(s, x, fy + Inches(0.35), fw, Inches(1.0), num, size=46, color=TEXT, bold=True,
              font=HEAD_FONT, align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.2), fy + Inches(1.5), fw - Inches(0.4), Inches(0.9), cap,
              size=12.5, color=MUTED, align=PP_ALIGN.CENTER, line_spacing=1.25)

add_text(s, MARGIN, Inches(5.55), CONTENT_W, Inches(0.9),
          "Properties this repository tests at the scale of its own test suite and deployment — not yet at a customer's.",
          size=13, color=SECONDARY, italic=True, align=PP_ALIGN.CENTER)

add_footer(s, 8)

# =============================================================================
# SLIDE 9 — What is not true yet
# =============================================================================
s = add_slide()
kicker(s, "Honesty, as a strength")
headline(s, "What is not true yet.", Inches(1.1), size=38, color=TEXT)
add_text(s, MARGIN, Inches(2.0), Inches(11.6), Inches(0.6),
          "Stated here rather than discovered later — because the buyers this is built for will ask.",
          size=14, color=SECONDARY, italic=True)

honest_items = [
    "No SOC 2, ISO 27001, or HIPAA certification.",
    "No support SLA.",
    "Human red-team engagement outstanding — not substitutable by the adversarial suite.",
    "Single node: no horizontal scaling, no failover.",
    "Pre-release and invite-only.",
]
hy = Inches(2.75)
row_h = Inches(0.72)
for i, item in enumerate(honest_items):
    y = hy + i * (row_h + Inches(0.05))
    rounded_rect(s, MARGIN, y, CONTENT_W, row_h, fill=SURFACE, line_color=BORDER, radius=0.12)
    thin_rule(s, MARGIN, y, Inches(0.06), color=ACCENT2, height=row_h)
    add_text(s, MARGIN + Inches(0.3), y, CONTENT_W - Inches(0.6), row_h, item,
              size=15.5, color=TEXT, anchor=MSO_ANCHOR.MIDDLE, font=HEAD_FONT)

add_footer(s, 9)

# =============================================================================
# SLIDE 10 — Go-to-market
# =============================================================================
s = add_slide()
kicker(s, "Go-to-market")
headline(s, "Land. Charge. Expand.", Inches(1.05), size=34)

gtm_y = Inches(1.95)
gtm_h = Inches(1.55)
gtm_w = CONTENT_W
card(s, MARGIN, gtm_y, gtm_w, gtm_h,
     "Land — with platform and security teams in regulated industries",
     "Usually through the SDK: they have an agent they cannot ship because of where it would run, "
     "and Abhed is the shortest path to shipping it.",
     title_size=16, body_size=13, accent_bar=ACCENT)

y2 = gtm_y + gtm_h + Inches(0.18)
h2 = Inches(1.55)
c2 = card(s, MARGIN, y2, gtm_w, h2,
     "Charge — per deployment, annually, by tier  [PROPOSED]",
     "Hosted, private, air-gapped, with support. The tiers are the same product; what differs is what "
     "the customer needs proven and who is on the hook for it. Prices are not announced.",
     title_size=16, body_size=13, accent_bar=ACCENT2)
pill(s, MARGIN + gtm_w - Inches(1.55), y2 + Inches(0.15), Inches(1.35), Inches(0.32),
     "PROPOSED", fill=RGBColor(0x2A, 0x1A, 0x10), text_color=AMBER, size=9.5)

y3 = y2 + h2 + Inches(0.18)
h3 = Inches(2.05)
card(s, MARGIN, y3, gtm_w, h3,
     "Expand — with further products on the same infrastructure layer  [ALL PLANNED]",
     "GPU & workload benchmarking · AI security & red team · observability for agents · "
     "inference infrastructure · private cloud deployment — each sharing the store, the identity, "
     "the policy, and the record.",
     title_size=16, body_size=13, accent_bar=GREEN)

add_footer(s, 10)

# =============================================================================
# SLIDE 11 — Roadmap / where we are
# =============================================================================
s = add_slide()
kicker(s, "Roadmap")
headline(s, "Where we are.", Inches(1.05), size=34)

col_w11 = Inches(5.6)
gap11 = Inches(0.5)
lx = MARGIN
rx = MARGIN + col_w11 + gap11
ty11 = Inches(1.95)
th11 = Inches(4.85)

# Shipped column
rounded_rect(s, lx, ty11, col_w11, th11, fill=SURFACE, line_color=GREEN, radius=0.04, line_w=Pt(1.3))
add_text(s, lx + Inches(0.25), ty11 + Inches(0.2), col_w11 - Inches(0.5), Inches(0.4),
          "SHIPPED", size=13, color=GREEN, bold=True, font=MONO_FONT)
shipped = [
    "Harness (CLI, headless, JSON, RPC, server)",
    "Console (Abhed Chat, reference app)",
    "Go SDK",
    "Structured output",
    "Parallel subagents with worktree isolation",
    "Scheduled runs",
    "OpenTelemetry export",
    "Access dashboard",
]
sy = ty11 + Inches(0.75)
for item in shipped:
    d = slide_dot = s.shapes.add_shape(MSO_SHAPE.OVAL, lx + Inches(0.28), sy + Inches(0.09), Inches(0.09), Inches(0.09))
    d.fill.solid(); d.fill.fore_color.rgb = GREEN; d.line.fill.background(); d.shadow.inherit = False
    add_text(s, lx + Inches(0.5), sy, col_w11 - Inches(0.75), Inches(0.42), item,
              size=13, color=TEXT, font=HEAD_FONT)
    sy += Inches(0.49)

# Next column
rounded_rect(s, rx, ty11, col_w11, th11, fill=SURFACE, line_color=BORDER, radius=0.04)
add_text(s, rx + Inches(0.25), ty11 + Inches(0.2), col_w11 - Inches(0.5), Inches(0.4),
          "NEXT", size=13, color=MUTED, bold=True, font=MONO_FONT)
nxt = [
    "Customer pilots",
    "Human red-team engagement",
    "Certifications path",
]
ny = ty11 + Inches(0.75)
for item in nxt:
    d = s.shapes.add_shape(MSO_SHAPE.OVAL, rx + Inches(0.28), ny + Inches(0.09), Inches(0.09), Inches(0.09))
    d.fill.solid(); d.fill.fore_color.rgb = MUTED; d.line.fill.background(); d.shadow.inherit = False
    add_text(s, rx + Inches(0.5), ny, col_w11 - Inches(0.75), Inches(0.42), item,
              size=13, color=SECONDARY, font=HEAD_FONT)
    ny += Inches(0.49)

add_text(s, rx + Inches(0.28), ty11 + th11 - Inches(0.85), col_w11 - Inches(0.55), Inches(0.7),
          "No dates promised beyond “next.”", size=11.5, color=MUTED, italic=True, font=HEAD_FONT)

add_footer(s, 11)

# =============================================================================
# SLIDE 12 — Ask / contact
# =============================================================================
s = add_slide()

thin_rule(s, MARGIN, Inches(1.85), Inches(1.0), color=ACCENT, height=Pt(3))
headline(s, "See it run on your infrastructure.", Inches(2.15), size=42, color=TEXT)

add_text(s, MARGIN, Inches(3.75), Inches(11), Inches(0.5),
          "support@zybuu.com", size=20, color=ACCENT, bold=True, font=MONO_FONT)
add_text(s, MARGIN, Inches(4.35), Inches(11), Inches(0.5),
          "zybuu.com", size=17, color=SECONDARY, font=MONO_FONT)
add_text(s, MARGIN, Inches(4.85), Inches(11), Inches(0.5),
          "abhed.zybuu.com/docs", size=17, color=SECONDARY, font=MONO_FONT)

add_footer(s, 12)

# ---------------------------------------------------------------- save -----
OUT = os.environ.get("DECK_OUT", os.path.join(os.path.dirname(os.path.abspath(__file__)), "abhed-deck.pptx"))
prs.save(OUT)
print(f"Saved {OUT}")
print(f"Slide count: {len(prs.slides._sldIdLst)}")
